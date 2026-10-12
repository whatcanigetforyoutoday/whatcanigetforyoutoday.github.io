# export-sheet.mjs가 부르는 엑셀 작성기: python write_xlsx.py rows.json out.xlsx
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

rows_path, out_path = sys.argv[1], sys.argv[2]
data = json.load(open(rows_path, encoding='utf-8'))

FONT = '맑은 고딕'
f_body = Font(name=FONT, size=10)
f_ref = Font(name=FONT, size=10, color='808080')
f_head = Font(name=FONT, size=10, bold=True, color='FFFFFF')
f_link = Font(name=FONT, size=10, color='1F4FD8', underline='single')
fill_edit_head = PatternFill('solid', fgColor='2F6BFF')
fill_ref_head = PatternFill('solid', fgColor='7A8699')
fill_edit = PatternFill('solid', fgColor='FFFBE6')
thin = Side(style='thin', color='D9DEE7')
border = Border(bottom=thin)
center = Alignment(horizontal='center', vertical='center')

MEAL_CATS = ['한식', '중식', '일식', '양식', '아시안', '분식', '고기', '버거·치킨', '샐러드', '기타']
CAFE_CATS = ['커피', '디저트', '베이커리', '아이스크림·빙수', '티']

wb = Workbook()

# ---------- 안내 ----------
g = wb.active
g.title = '안내'
g.sheet_view.showGridLines = False
g.column_dimensions['A'].width = 26
g.column_dimensions['B'].width = 70
lines = [
    ('오늘 뭐 먹지 · 식당 목록 수정 방법', None, 'title'),
    (None, None, None),
    ('1. 고치는 칸', "'식당'·'카페' 시트에서 파란 머리글(노란 바탕) 칸만 고쳐 주세요. 회색 칸은 참고용이에요.", None),
    ('2. ID', '맨 앞 ID는 가게를 구분하는 번호라 절대 고치거나 지우지 마세요.', None),
    ('3. 가게 빼기', "'목록에서 빼기' 칸에 O를 적으면 앱에서 아예 안 나와요. 행을 지우지 말고 O로 표시해 주세요.", None),
    ('4. 인당 2.5만원↑', 'O = 비싼 곳(앱의 "인당 2.5만원 넘는 곳 빼기"가 켜져 있으면 제외), X = 아님. 처음 값은 이름·종류로 자동 추정한 거예요.', None),
    ('5. 영업', '언제나 / 저녁만(점심엔 제외) / 점심만(저녁엔 제외). 여는·닫는 시간을 적으면 그 시간이 우선이에요.', None),
    ('6. 시간 적는 법', '24시간제 글자로: 11:30, 21:00. 새벽에 닫으면 02:00처럼 적으면 돼요.', None),
    ('7. 쉬는 요일', '쉼표로: 일  또는  토,일', None),
    ('8. 메모', '대표메뉴·가격·특징을 자유롭게. 그 가게가 뽑히면 결과 화면에 보여요.', None),
    ('9. 새 가게 추가', '맨 아래에 이름·종류·주소를 적고 ID 칸은 비워 두세요. 위치는 주소로 찾아 넣어 드려요.', None),
    ('10. 다 고쳤으면', '이 파일을 저장해서 그대로 보내 주세요. 앱에 반영해 드릴게요.', None),
    (None, None, None),
    ('예시 (한 줄)', '이름: 담솥 압구정점 · 종류: 한식 · 인당 2.5만원↑: X · 영업: 언제나 · 여는 시간: 11:00 · 닫는 시간: 21:30 · 쉬는 요일: 일 · 메모: 장어솥밥 19,000원, 12시 넘으면 웨이팅', None),
    (None, None, None),
    ('도보 거리 기준 (앱과 같게)', None, 'sub'),
]
for i, (a, b, kind) in enumerate(lines, start=1):
    if a:
        c = g.cell(i, 1, a)
        c.font = Font(name=FONT, size=16 if kind == 'title' else 10, bold=kind in ('title', 'sub') or bool(b))
    if b:
        c = g.cell(i, 2, b)
        c.font = f_body
        c.alignment = Alignment(wrap_text=True, vertical='top')
r0 = len(lines) + 1
params = [
    ('직선→도보 보정 배율', 1.5, '지도 직선거리 × 이 값 = 예상 도보거리. 앱(app.html의 DETOUR)과 같은 값'),
    ('5분 이내 (m)', 400, '사용자가 정한 기준 (2026-10-12)'),
    ('7분 이내 (m)', 560, '사용자가 정한 기준 (2026-10-12)'),
    ('10분 이내 (m)', 800, '사용자가 정한 기준 (2026-10-12)'),
]
for k, (label, val, note) in enumerate(params):
    g.cell(r0 + k, 1, label).font = f_body
    c = g.cell(r0 + k, 2, val)
    c.font = Font(name=FONT, size=10, color='0000FF')
    c.alignment = Alignment(horizontal='left')
    c.comment = Comment(note, '오늘 뭐 먹지')
P_DETOUR, P5, P7, P10 = (f"'안내'!$B${r0 + k}" for k in range(4))


def build(ws, rows, cols, cats):
    ws.freeze_panes = 'C2'
    ws.sheet_view.zoomScale = 100
    for j, (key, head, width, editable) in enumerate(cols, start=1):
        c = ws.cell(1, j, head)
        c.font = f_head
        c.fill = fill_edit_head if editable else fill_ref_head
        c.alignment = center
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 22
    letter = {key: ws.cell(1, j).column_letter for j, (key, *_rest) in enumerate(cols, start=1)}
    for i, r in enumerate(rows, start=2):
        for j, (key, head, width, editable) in enumerate(cols, start=1):
            if key == 'walk':
                v = f'=ROUND({letter["dist"]}{i}*{P_DETOUR},-1)'
            elif key == 'tier':
                w = f'{letter["walk"]}{i}'
                v = f'=IF({w}<={P5},"5분 이내",IF({w}<={P7},"7분 이내",IF({w}<={P10},"10분 이내","10분 밖")))'
            elif key == 'kakao':
                v = '열기' if r['kakao'] else ''
            else:
                v = r.get(key, '')
            c = ws.cell(i, j, v)
            c.font = f_body if editable else f_ref
            c.border = border
            if editable:
                c.fill = fill_edit
            if key in ('pricey', 'when', 'open', 'close', 'hide', 'tier', 'cat'):
                c.alignment = center
            if key == 'kakao' and r['kakao']:
                c.hyperlink = r['kakao']
                c.font = f_link
            if key in ('dist', 'walk'):
                c.number_format = '#,##0'
            if key in ('open', 'close'):
                c.number_format = '@'
    last = len(rows) + 1 + 200   # 새 가게를 적을 여유 줄까지
    ws.auto_filter.ref = f'A1:{ws.cell(1, len(cols)).column_letter}{len(rows) + 1}'

    def dv(key, options, prompt):
        d = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=True, showErrorMessage=True,
                           errorTitle='값 확인', error=prompt)
        ws.add_data_validation(d)
        d.add(f'{letter[key]}2:{letter[key]}{last}')
    dv('cat', cats, '목록에 있는 종류 중에서 골라 주세요')
    if 'pricey' in letter:
        dv('pricey', ['O', 'X'], 'O 또는 X로 적어 주세요')
    dv('when', ['언제나', '저녁만', '점심만'], '언제나 / 저녁만 / 점심만 중에서 골라 주세요')
    dv('hide', ['O'], '빼려면 O, 아니면 비워 두세요')
    # 머리글 설명
    tips = {'id': '가게 구분 번호. 고치거나 지우지 마세요', 'pricey': 'O = 인당 2.5만원 넘음', 'when': '시간을 적으면 시간이 우선',
            'hide': 'O를 적으면 앱에서 안 나와요', 'walk': '직선거리 × 안내 시트의 보정 배율', 'tier': '안내 시트의 도보 기준으로 자동 계산'}
    for key, tip in tips.items():
        if key in letter:
            ws[f'{letter[key]}1'].comment = Comment(tip, '오늘 뭐 먹지')


COMMON_TAIL = [
    ('kakaoCat', '카카오 분류', 22, False), ('address', '주소', 30, False),
    ('dist', '직선거리(m)', 11, False), ('walk', '예상 도보거리(m)', 14, False), ('tier', '도보 구분', 11, False),
    ('kakao', '카카오맵', 9, False),
]
meal_cols = [('id', 'ID', 13, False), ('name', '이름', 26, True), ('cat', '종류', 10, True), ('pricey', '인당 2.5만원↑', 12, True),
             ('when', '영업', 9, True), ('open', '여는 시간', 9, True), ('close', '닫는 시간', 9, True), ('off', '쉬는 요일', 10, True),
             ('hide', '목록에서 빼기', 12, True), ('note', '메모 (대표메뉴 등)', 36, True)] + COMMON_TAIL
cafe_cols = [c for c in meal_cols if c[0] != 'pricey']

build(wb.create_sheet('식당'), data['meal'], meal_cols, MEAL_CATS)
build(wb.create_sheet('카페'), data['cafe'], cafe_cols, CAFE_CATS)
wb.active = 1
wb.calculation.fullCalcOnLoad = True   # 엑셀이 열 때 수식을 계산하도록
wb.save(out_path)
