# 사용자가 고친 엑셀을 앱에 반영: python tools/import_xlsx.py 식당목록.xlsx
# → data/overrides.json (앱이 data/restaurants.json·cafes.json 위에 덮어써서 씀)
# ID가 빈 새 가게는 주소로 위치를 찾아 추가해요 (kakao_key.txt 필요)
import json, sys, datetime, urllib.request, urllib.parse, hashlib, math
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
src = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / '식당목록.xlsx'
wb = load_workbook(src, data_only=True)

DAY = '일월화수목금토'
WHEN = {'언제나': 'all', '저녁만': 'dinner', '점심만': 'lunch'}
company = json.loads((ROOT / 'data/company.json').read_text(encoding='utf-8'))


def text(v):
    return '' if v is None else str(v).strip()


def hhmm(v):
    if isinstance(v, datetime.time):
        return f'{v.hour:02d}:{v.minute:02d}'
    if isinstance(v, datetime.datetime):
        return f'{v.hour:02d}:{v.minute:02d}'
    if isinstance(v, (int, float)) and 0 <= v < 1:          # 엑셀이 시간으로 바꿔 저장한 경우
        m = round(v * 24 * 60)
        return f'{m // 60:02d}:{m % 60:02d}'
    s = text(v).replace('시', ':').replace(' ', '')
    if not s:
        return ''
    parts = s.split(':')
    try:
        h, m = int(parts[0]), int(parts[1] or 0) if len(parts) > 1 else 0
        return f'{h:02d}:{m:02d}'
    except ValueError:
        raise ValueError(f'시간 형식을 알 수 없어요: {v}')


def days(v):
    return sorted({DAY.index(ch) for ch in text(v) if ch in DAY})


def dist_m(a, b):
    R, rad = 6371000, math.radians
    dlat, dlng = rad(b['lat'] - a['lat']), rad(b['lng'] - a['lng'])
    h = math.sin(dlat / 2) ** 2 + math.cos(rad(a['lat'])) * math.cos(rad(b['lat'])) * math.sin(dlng / 2) ** 2
    return round(2 * R * math.asin(math.sqrt(h)))


def geocode(addr):
    key = (ROOT / 'kakao_key.txt').read_text(encoding='utf-8').strip()
    for path in ('address', 'keyword'):
        url = f'https://dapi.kakao.com/v2/local/search/{path}.json?' + urllib.parse.urlencode({'query': addr})
        req = urllib.request.Request(url, headers={'Authorization': f'KakaoAK {key}'})
        docs = json.load(urllib.request.urlopen(req, timeout=10))['documents']
        if docs:
            return round(float(docs[0]['y']), 6), round(float(docs[0]['x']), 6)
    return None, None


out = {'meal': {}, 'cafe': {}, 'added': {'meal': [], 'cafe': []}}
problems, counts = [], {}
for sheet, kind in (('식당', 'meal'), ('카페', 'cafe')):
    ws = wb[sheet]
    head = [text(c.value) for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    n_edit = n_hide = n_new = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        get = lambda h: row[col[h]] if h in col and col[h] < len(row) else None
        rid, name = text(get('ID')), text(get('이름'))
        if not rid and not name:
            continue
        try:
            o = {'name': name, 'cat': text(get('종류'))}
            if kind == 'meal':
                o['pricey'] = text(get('인당 2.5만원↑')).upper() == 'O'
            # 영업 칸이 없거나 비어 있으면 적지 않아요 → 앱이 이름·종류로 자동 판단
            if text(get('영업')) in WHEN:
                o['when'] = WHEN[text(get('영업'))]
            o['open'], o['close'] = hhmm(get('여는 시간')), hhmm(get('닫는 시간'))
            if bool(o['open']) != bool(o['close']):
                raise ValueError('여는 시간과 닫는 시간을 둘 다 적어야 해요')
            o['off'] = days(get('쉬는 요일'))
            o['hide'] = text(get('목록에서 빼기')).upper() == 'O'
            o['note'] = text(get('메모 (대표메뉴 등)'))
        except ValueError as e:
            problems.append(f'[{sheet}] {name}: {e}')
            continue
        o = {k: v for k, v in o.items() if (v not in ('', [], None) and v is not False) or k == 'pricey'}
        n_hide += o.get('hide', False)
        if rid:
            out[kind][rid] = o
            n_edit += 1
        else:
            addr = text(get('주소'))
            lat, lng = geocode(addr or name) if (addr or name) else (None, None)
            if lat is None:
                problems.append(f'[{sheet}] 새 가게 {name}: 주소로 위치를 못 찾았어요 ({addr})')
                continue
            nid = 'x' + hashlib.md5((name + addr).encode()).hexdigest()[:10]
            item = {'id': nid, 'name': name, 'cat': o.get('cat') or '기타', 'address': addr, 'lat': lat, 'lng': lng,
                    'dist': dist_m(company, {'lat': lat, 'lng': lng}), 'kakao': '', 'naver': '', 'menus': [], 'note': ''}
            out['added'][kind].append(item)
            out[kind][nid] = o
            n_new += 1
    # 엑셀에서 행을 지운 가게는 '목록에서 빼기'로 봐요 (나중에 행을 다시 넣으면 살아나요)
    base = json.loads((ROOT / ('data/restaurants.json' if kind == 'meal' else 'data/cafes.json')).read_text(encoding='utf-8'))
    for r in base:
        if r['id'] not in out[kind]:
            out[kind][r['id']] = {'hide': True}
            n_hide += 1
    counts[sheet] = (n_edit, n_hide, n_new)

(ROOT / 'data/overrides.json').write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding='utf-8')
for sheet, (e, h, n) in counts.items():
    print(f'{sheet}: {e}곳 반영 · 빼기 {h}곳 · 새 가게 {n}곳')
for p in problems:
    print('확인 필요 →', p)
